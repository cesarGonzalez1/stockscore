import math, pandas as pd, pytest
from stockscore.config import load_config
from stockscore.models import Snapshot, cagr
from stockscore.providers import get_provider
from stockscore.metrics.quality import piotroski, altman
from stockscore.metrics.valuation import project_pv, implied_growth
from stockscore.metrics.risk import swing_probability
from stockscore.metrics.trend import rsi_wilder, slope_tstat
from stockscore.scoring import metric_score, score_company, score_universe, combine, score_from_metrics
from stockscore.metrics import compute_all
from stockscore.ranking import to_frame, top_n
from stockscore.storage import Store
from stockscore.synthetic import synthetic_universe
from stockscore import validate as V
from stockscore.providers.yfinance_provider import build_fiscal, pick

CFG = load_config()
import pathlib
FIX = str(pathlib.Path(__file__).resolve().parent / "fixtures")
P = get_provider("snapshot", directory=FIX)   # fixtures propios: no dependen de lo que bajes con fetch


# ---------- datos reales: identidades contables ----------
@pytest.mark.parametrize("t", ["AAPL", "NKE"])
def test_balance_identity(t):
    for f in P.get(t).fiscal:
        assert abs(f["total_assets"] - f["total_liabilities"] - f["equity"]) < 2


# ---------- piotroski / altman ----------
def test_piotroski_matches_hand_calc():
    assert piotroski(P.get("AAPL"))[0] == 8          # falla solo CFO>NI
    p, sig = piotroski(P.get("AAPL")); assert sig["CFO>NI"] is False
    assert piotroski(P.get("NKE"))[0] == 4           # ROA>0, CFO>0, menos deuda LP, margen bruto sube


def test_altman_z2_formula():
    z, z2 = altman(P.get("AAPL"))
    f = P.get("AAPL").fy(0)
    exp = 6.56 * (f["current_assets"] - f["current_liabilities"]) / f["total_assets"] + 3.26 * f["retained_earnings"] / f["total_assets"] \
        + 6.72 * f["operating_income"] / f["total_assets"] + 1.05 * f["equity"] / f["total_liabilities"]
    assert z2 == pytest.approx(exp)
    assert z > 10   # original con capitalización actual enorme


# ---------- DCF ----------
def test_reverse_dcf_roundtrip():
    for g in (-0.05, 0.03, 0.12, 0.25):
        v = project_pv(100, g, 0.09, 0.03, 10)
        assert implied_growth(100, v, 0.09, 0.03, 10) == pytest.approx(g, abs=1e-6)


def test_reverse_dcf_none_when_fcf_negative():
    assert implied_growth(-5, 100, 0.09, 0.03, 10) is None


def test_pv_monotonic_in_growth():
    assert project_pv(100, 0.10, .09, .03, 10) > project_pv(100, 0.05, .09, .03, 10)


# ---------- riesgo / tendencia ----------
def test_swing_probability_bounds_and_driftless_limit():
    pu, pd_, pn = swing_probability(0.3, 0.0, 0.08, 0.05, 2000, sims=3000, seed=1)   # horizonte largo
    assert pu + pd_ + pn == pytest.approx(1)
    assert pu == pytest.approx(5 / 13, abs=0.05)     # random walk sin deriva: stop/(stop+target)
    assert swing_probability(0.3, 0.0, 0.08, 0.05, 20, seed=1)[2] > swing_probability(0.3, 0.0, 0.08, 0.05, 200, seed=1)[2]


def test_rsi_extremes():
    up = pd.Series(range(1, 40), dtype=float)
    assert rsi_wilder(up) == 100
    assert rsi_wilder(up[::-1].reset_index(drop=True)) < 1


def test_slope_t_positive_on_uptrend():
    s = pd.Series([100 * 1.001 ** i for i in range(80)])
    assert slope_tstat(s) > 50


# ---------- scoring ----------
def test_metric_score_mapping():
    sp = {"low": 0, "high": 10, "higher_better": True}
    assert metric_score(5, sp) == 50 and metric_score(-3, sp) == 0 and metric_score(99, sp) == 100 and metric_score(None, sp) is None
    assert metric_score(2, {**sp, "higher_better": False}) == 80


def test_missing_metric_renormalizes_not_zero():
    m, n = compute_all(P.get("AAPL"), CFG, P.get("SPY"))
    full = score_from_metrics(P.get("AAPL"), m, n, CFG)
    m2 = {k: v for k, v in m.items() if k not in ("peg", "dcf_mos")}
    part = score_from_metrics(P.get("AAPL"), m2, n, CFG)
    assert part.pillars["value"]["coverage"] < 1
    assert abs(part.total - full.total) < 10        # no se desploma por dato faltante


def test_combine_weights():
    pil = {"a": {"score": 100, "coverage": 1}, "b": {"score": 0, "coverage": 1}}
    assert combine(pil, {"a": 3, "b": 1})[0] == 75
    assert combine({"a": {"score": 100, "coverage": 1}, "b": {"score": None, "coverage": 0}}, {"a": 1, "b": 1})[0] == 100


def test_veto_caps_score():
    s = P.get("AAPL")
    s.fiscal[0] = {**s.fiscal[0], "retained_earnings": -5e6, "equity": -1e5}
    r = score_company(s, CFG, "swing", P.get("SPY"))
    assert r.vetoes and r.total <= CFG["vetoes"]["veto_cap"] and r.signal == "EVITAR"


def test_real_ranking_apple_above_nike_both_profiles():
    for prof in ("swing", "investor"):
        r = score_universe([P.get("NKE"), P.get("AAPL")], CFG, prof, P.get("SPY"))
        assert [x.ticker for x in r] == ["AAPL", "NKE"]
        assert all(0 <= x.total <= 100 for x in r)


# ---------- ranking / storage (50 empresas) ----------
@pytest.fixture(scope="module")
def universe50():
    snaps = [P.get("AAPL"), P.get("NKE")] + synthetic_universe(48, seed=1)
    return snaps, score_universe(snaps, CFG, "swing", P.get("SPY"))


def test_top10_of_50(universe50):
    snaps, res = universe50
    assert len(res) == 50
    top = top_n(res, 10)
    assert len(top) == 10 and list(top["rank"]) == list(range(1, 11))
    assert top["total"].is_monotonic_decreasing
    assert len(to_frame(res)) == 50


def test_storage_roundtrip(tmp_path, universe50):
    _, res = universe50
    st = Store(str(tmp_path / "t.db"))
    rid = st.save_run(res, "swing", "h", "x")
    rk = st.ranking(rid, n=10)
    assert len(rk) == 10 and rk.loc[0, "total"] == max(r.total for r in res)
    d = st.get("AAPL", rid)
    assert d["ticker"] == "AAPL" and "piotroski" in d["metrics"] and d["rank"] >= 1
    st.save_run(res, "swing", "h", "y")
    assert len(st.history("AAPL")) == 2
    assert st.get("ZZZZ") is None


def test_sector_filter(tmp_path, universe50):
    _, res = universe50
    st = Store(str(tmp_path / "t.db")); st.save_run(res, "swing")
    df = st.ranking(sector="Technology")
    assert set(df.sector) == {"Technology"}


# ---------- validación ----------
def test_validation_outputs(universe50):
    snaps, res = universe50
    b = V.benchmarks(snaps, CFG, P.get("SPY"))
    c = V.correlations(res, b)
    assert len(c) == 5 and c["spearman"].between(-1, 1).all()
    st = V.weight_stability(snaps, CFG, n_runs=30, bench=P.get("SPY"))
    assert 0 <= st["topn_kept_mean"] <= 1 and st["spearman_mean"] > 0.8
    rc = V.reported_check([P.get("AAPL"), P.get("NKE")], CFG)
    assert rc.loc[rc.ticker == "AAPL", "piotroski_calc"].iloc[0] == 8


def test_synthetic_universe_valid():
    for s in synthetic_universe(20, seed=3):
        f = s.fy(0)
        assert abs(f["total_assets"] - f["total_liabilities"] - f["equity"]) < 1e-6
        assert len(s.closes()) >= 250


# ---------- proveedor yfinance (con datos simulados, sin red) ----------
def test_yfinance_mapping_with_fake_frames():
    col = pd.Timestamp("2025-09-30")
    inc = pd.DataFrame({col: [100e6, 40e6, 20e6, 15e6]}, index=["Total Revenue", "Gross Profit", "Operating Income", "Net Income"])
    bal = pd.DataFrame({col: [200e6, 80e6, 50e6, 120e6, 80e6, 5e6]}, index=["Total Assets", "Current Assets", "Current Liabilities", "Total Liabilities Net Minority Interest", "Stockholders Equity", "Retained Earnings"])
    cf = pd.DataFrame({col: [25e6, -6e6]}, index=["Operating Cash Flow", "Capital Expenditure"])
    f = build_fiscal(inc, bal, cf)[0]
    assert f["revenue"] == 100 and f["capex"] == 6 and f["cfo"] == 25 and f["total_assets"] == 200
    assert pick(inc, ["Nope"], col) is None


def test_excel_export(tmp_path, universe50):
    import openpyxl
    from stockscore.export import to_excel
    _, res = universe50
    p = tmp_path / "r.xlsx"
    to_excel(res, str(p), {"Extra": pd.DataFrame({"a": [1]})})
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Ranking", "Metricas_crudas", "Puntajes_metricas", "Alertas", "Extra"]
    assert wb["Ranking"].max_row == 51


def test_missing_total_liabilities_is_repaired_and_universe_survives():
    d = P.get("AAPL").d
    for f in d["fiscal"]:
        f["total_liabilities"] = 0.0          # como Yahoo cuando omite la fila
    s = Snapshot(d)
    assert s.fy(0)["total_liabilities"] == pytest.approx(s.fy(0)["total_assets"] - s.fy(0)["equity"])
    z, z2 = altman(s); assert z2 is not None
    r = score_universe([s, P.get("NKE")], CFG, "swing", P.get("SPY"))
    assert len(r) == 2


def test_zero_current_liabilities_and_assets_do_not_crash():
    d = P.get("NKE").d
    d["fiscal"][0]["current_liabilities"] = 0.0
    d["fiscal"][0]["equity"] = 0.0
    d["fiscal"][0]["total_liabilities"] = 0.0
    r = score_universe([Snapshot(d)], CFG, "swing", P.get("SPY"))
    assert len(r) == 1


def test_bad_ticker_is_skipped_not_fatal():
    d = P.get("AAPL").d
    d["fiscal"] = [{"year": 2025}]            # datos rotos
    r = score_universe([Snapshot(d), P.get("NKE")], CFG, "swing", P.get("SPY"))
    assert [x.ticker for x in r] == ["NKE"]


def test_dcf_growth_fallback_uses_forward_not_only_history():
    d = P.get("AAPL").d
    d["market"]["eps_growth_3y"] = None; d["market"]["revenue_growth_3y"] = None   # como Yahoo
    m, _ = compute_all(Snapshot(d), CFG, P.get("SPY"))
    hist = m["rev_cagr_3y"]
    assert m["base_growth"] > hist + 0.03          # ya no queda pegado al CAGR histórico (1.8%)
    assert m["base_growth"] <= CFG["assumptions"]["growth_cap"]
